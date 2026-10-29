#!/usr/bin/env python3
"""Build a single Word document containing every HireAssist ADR."""
import re, glob, os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

REPO = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
OUT = os.path.join(REPO, "docs", "HireAssist-ADRs.docx")
INDEX = os.path.join(REPO, "docs", "adr", "INDEX.md")

INK = RGBColor(0x1A, 0x1A, 0x1A)
ACCENT = RGBColor(0x1B, 0x5E, 0x8C)
MUTED = RGBColor(0x55, 0x5F, 0x6B)
CODEBG = "EEF2F6"

INLINE = re.compile(r'(\*\*.+?\*\*|~~.+?~~|`[^`]+`|\[[^\]]+\]\([^)]*\)|\*[^*\n]+?\*)')


def shade(el, fill):
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:fill'), fill)
    el.append(sh)


def add_runs(par, text, base_italic=False, base_color=None, base_bold=False, base_strike=False):
    """Render a subset of inline markdown into runs, recursing so that a code span
    inside bold (or bold inside a list item) is still rendered as code."""
    def emit(t, **kw):
        add_runs(par, t,
                 base_italic=kw.get('italic', base_italic),
                 base_color=base_color,
                 base_bold=kw.get('bold', base_bold),
                 base_strike=kw.get('strike', base_strike))

    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith('**') and tok.endswith('**') and len(tok) > 4:
            emit(tok[2:-2], bold=True); continue
        if tok.startswith('~~') and tok.endswith('~~') and len(tok) > 4:
            emit(tok[2:-2], strike=True); continue
        if tok.startswith('*') and tok.endswith('*') and len(tok) > 2:
            emit(tok[1:-1], italic=True); continue
        if tok.startswith('[') and '](' in tok:
            emit(tok[1:tok.index('](')]); continue

        if tok.startswith('`') and tok.endswith('`') and len(tok) > 2:
            r = par.add_run(tok[1:-1])
            r.font.name = 'Consolas'
            r.font.size = Pt(9.5)
            r.font.color.rgb = ACCENT
            shade(r._element.get_or_add_rPr(), CODEBG)
        else:
            r = par.add_run(tok)
            if base_color is not None:
                r.font.color.rgb = base_color
        if base_italic:
            r.italic = True
        if base_bold:
            r.bold = True
        if base_strike:
            r.font.strike = True


def is_table_sep(line):
    return bool(re.match(r'^\|[\s:\-\|]+\|$', line.strip()))


def split_row(line):
    cells = line.strip().strip('|').split('|')
    return [c.strip() for c in cells]


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    t = doc.add_table(rows=len(rows), cols=len(header))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for j, cell in enumerate(header):
        c = t.cell(0, j)
        c.text = ''
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        add_runs(p, cell)
        for r in p.runs:
            r.bold = True
            r.font.size = Pt(9.5)
        shade(c._tc.get_or_add_tcPr(), 'E8EEF4')
    for i, row in enumerate(body, start=1):
        for j in range(len(header)):
            c = t.cell(i, j)
            c.text = ''
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            add_runs(p, row[j] if j < len(row) else '')
            for r in p.runs:
                if r.font.size is None:
                    r.font.size = Pt(9.5)
    doc.add_paragraph()


def render(doc, md, first):
    lines = md.split('\n')
    i = 0
    para, quote, bullets = [], [], []

    def flush_para():
        nonlocal para
        if para:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(8)
            add_runs(p, ' '.join(para))
            para = []

    def flush_quote():
        nonlocal quote
        if quote:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.right_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(12)
            add_runs(p, ' '.join(quote), base_italic=True, base_color=MUTED)
            quote = []

    def flush_bullets():
        nonlocal bullets
        for style, txt in bullets:
            p = doc.add_paragraph(style=style)
            p.paragraph_format.space_after = Pt(4)
            add_runs(p, txt)
        bullets = []

    def flush_all():
        flush_para(); flush_quote(); flush_bullets()

    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        s = line.strip()

        if s.startswith('#'):
            flush_all()
            level = len(s) - len(s.lstrip('#'))
            text = s[level:].strip()
            if level == 1:
                if not first:
                    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                h = doc.add_heading(level=1)
                add_runs(h, text)
            else:
                h = doc.add_heading(level=min(level, 4))
                add_runs(h, text)
            i += 1
            continue

        if s.startswith('|') and s.endswith('|'):
            flush_all()
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                if not is_table_sep(lines[i]):
                    rows.append(split_row(lines[i]))
                i += 1
            if rows:
                add_table(doc, rows)
            continue

        if s.startswith('>'):
            flush_para(); flush_bullets()
            quote.append(s.lstrip('>').strip())
            i += 1
            continue

        if re.match(r'^[-*]\s+\S', s):
            flush_para(); flush_quote()
            txt = re.sub(r'^[-*]\s+', '', s)
            j = i + 1
            while j < len(lines) and lines[j].startswith('  ') and lines[j].strip() \
                    and not re.match(r'^\s*[-*]\s', lines[j]) and not lines[j].strip().startswith('|'):
                txt += ' ' + lines[j].strip()
                j += 1
            bullets.append(('List Bullet', txt))
            i = j
            continue

        if re.match(r'^\d+\.\s+\S', s):
            flush_para(); flush_quote()
            txt = re.sub(r'^\d+\.\s+', '', s)
            j = i + 1
            while j < len(lines) and lines[j].startswith('   ') and lines[j].strip() \
                    and not re.match(r'^\s*\d+\.\s', lines[j]) and not lines[j].strip().startswith('|'):
                txt += ' ' + lines[j].strip()
                j += 1
            bullets.append(('List Number', txt))
            i = j
            continue

        if s in ('---', '***', '___'):
            flush_all()
            i += 1
            continue

        if not s:
            flush_all()
            i += 1
            continue

        flush_quote(); flush_bullets()
        para.append(s)
        i += 1

    flush_all()


def latest_date(files):
    """The stack is current as of its newest record."""
    import datetime
    dates = []
    for f in files:
        m = re.search(r'\n\s*\*\*?Date:?\*\*?\s*(\d{4}-\d{2}-\d{2})', open(f).read())
        if not m:
            m = re.search(r'(\d{4}-\d{2}-\d{2})', open(f).read())
        if m:
            dates.append(m.group(1))
    if not dates:
        raise ValueError('no date found in any ADR')
    d = datetime.date.fromisoformat(max(dates))
    return f'{d.day} {d.strftime("%B %Y")}'


def index_table(path, heading):
    """Lift one markdown table out of INDEX.md, so the fact lives in exactly one file."""
    rows, seen = [], False
    for line in open(path).read().split('\n'):
        s = line.strip()
        if s.startswith('#') and heading in s:
            seen = True
            continue
        if not seen:
            continue
        if s.startswith('|') and s.endswith('|'):
            if re.match(r'^\|[\s:\-\|]+\|$', s):
                continue
            rows.append([c.strip() for c in s.strip('|').split('|')])
        elif rows:
            break
    if not rows or len(rows[0]) != 3:
        raise ValueError(f'table under "{heading}" not found in {path}, or its shape changed')
    return rows


def main():
    doc = Document()

    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    for name, size in (('Heading 1', 18), ('Heading 2', 14), ('Heading 3', 11.5), ('Heading 4', 10.5)):
        st = doc.styles[name]
        st.font.name = 'Calibri'
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = ACCENT if name == 'Heading 1' else INK
        st.paragraph_format.space_before = Pt(14 if name in ('Heading 1', 'Heading 2') else 10)
        st.paragraph_format.space_after = Pt(6)

    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(1.0)
        s.top_margin = s.bottom_margin = Inches(0.9)

    files = sorted(glob.glob(os.path.join(REPO, 'docs', 'adr', 'ADR-*.md')))

    # --- title page ---
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.paragraph_format.space_before = Pt(150)
    r = t.add_run('HireAssist')
    r.font.size = Pt(34); r.bold = True; r.font.color.rgb = ACCENT

    t2 = doc.add_paragraph()
    t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t2.add_run('Architecture Decision Records')
    r.font.size = Pt(17); r.font.color.rgb = MUTED

    t3 = doc.add_paragraph()
    t3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t3.paragraph_format.space_before = Pt(26)
    r = t3.add_run('Software Architecture term project · Chulalongkorn University')
    r.font.size = Pt(10.5); r.font.color.rgb = MUTED

    t4 = doc.add_paragraph()
    t4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t4.add_run('Patiphon Puntusin · Thanabul Parodom · Thanwarat Korcharoenkiat · Rerngrit Jangsri')
    r.font.size = Pt(10.5); r.font.color.rgb = MUTED

    t5 = doc.add_paragraph()
    t5.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t5.paragraph_format.space_before = Pt(20)
    r = t5.add_run(latest_date(files))
    r.font.size = Pt(10); r.font.color.rgb = MUTED


    # --- contents ---
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    h = doc.add_heading(level=1)
    h.add_run('Contents')
    rows = [['ID', 'Decision', 'Status']]
    for f in files:
        text = open(f).read()
        title = text.split('\n', 1)[0].lstrip('# ').strip()
        num, _, rest = title.partition(':')
        m = re.search(r'### Status\s*\n+\s*\*\*(.+?)\*\*', text)
        status = m.group(1).strip().rstrip('.') if m else ''
        rows.append([num.strip(), rest.strip(), status])
    add_table(doc, rows)

    p = doc.add_paragraph()
    add_runs(p, 'These records are a stack. Each states what was decided on its date and nothing '
                'after it; an earlier record is never edited because of a later one. To find the '
                'current position on a question, read forward from the record that first raised it.')
    for r in p.runs:
        r.font.color.rgb = MUTED
        r.italic = True

    # --- what later records have changed ---
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    h = doc.add_heading(level=1)
    h.add_run('What later records have changed')
    p = doc.add_paragraph()
    add_runs(p, 'Because an earlier record is never edited, this is the only place that says '
                'which parts of an earlier one no longer hold. Read it before acting on any '
                'record below.')
    add_table(doc, index_table(INDEX, 'What later records have changed'))

    # --- the records ---
    for f in files:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        render(doc, open(f).read(), first=True)

    cp = doc.core_properties
    cp.title = 'HireAssist — Architecture Decision Records'
    cp.subject = 'Software Architecture term project, Chulalongkorn University'
    cp.author = 'Patiphon Puntusin, Thanabul Parodom, Thanwarat Korcharoenkiat, Rerngrit Jangsri'

    doc.save(OUT)
    print('wrote', OUT)
    print('records:', len(files))
    for f in files:
        print('  -', os.path.basename(f))


if __name__ == '__main__':
    main()
